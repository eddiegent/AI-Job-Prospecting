"""Regression tests for scripts/claims_guard.py (anti-exaggeration guard)."""
from __future__ import annotations

from scripts.claims_guard import (
    _norm,
    calibration_bullets,
    check,
    qualified_skills_from_rows,
    stack_hint_techs,
)

MASTER = _norm("Python (lecture/adaptation de scripts)\nC#, WinForms, WPF\nstand-ups quotidiens\nSQL Server, PostgreSQL\nGitLab, GitHub")
FAMILIAR = ["Docker", "RabbitMQ", "CI/CD", "ASP .Net Core"]
# Skills-table rows as the master-CV docx reader sees them (heading | items).
ROWS = [
    ["Langages", "C#, Python (lecture/adaptation de scripts)"],
    ["Architecture", "Clean Architecture (usage pragmatique), SOLID"],
]
QUALIFIED = qualified_skills_from_rows(ROWS)


def run(cv=None, letter=None, offer=None, addendum="", profile=None):
    docs = {"tailored_cv.json": cv or {}, "letter.json": letter or {}}
    return check(None, MASTER, FAMILIAR, addendum, docs, offer or {},
                 qualified=QUALIFIED, profile=profile)


def test_intensity_wording_is_flagged():
    errs = run(cv={"summary_paragraphs": ["Pratique quotidienne de WCF"]})
    assert any(e.startswith("[intensity]") for e in errs)


def test_intensity_wording_present_in_source_is_allowed():
    assert run(cv={"experience": [{"bullets": ["Mise en place de stand-ups quotidiens"]}]}) == []


def test_familiar_skill_outside_familiar_section_is_flagged():
    cv = {"skills_sections": [{"heading": "Outils", "items": ["Docker"]}]}
    assert any(e.startswith("[familiar]") for e in run(cv=cv))


def test_familiar_skill_in_familiar_section_is_ok():
    cv = {"skills_sections": [{"heading": "Familier / à approfondir", "items": ["Docker", "RabbitMQ"]}]}
    assert run(cv=cv) == []


def test_unsourced_skill_flagged_and_sourced_ok():
    cv = {"skills_sections": [{"heading": "Langages", "items": ["C#", "Oracle"]}]}
    errs = run(cv=cv)
    assert len(errs) == 1 and "Oracle" in errs[0]


def test_python_must_keep_qualifier():
    cv = {"skills_sections": [{"heading": "Langages", "items": ["Python"]}]}
    errs = run(cv=cv)
    assert errs == ["[calibration] tailored_cv.json: 'Python' must keep the master qualifier '(lecture/adaptation de scripts)'"]
    cv = {"skills_sections": [{"heading": "Langages", "items": ["Python (lecture/adaptation de scripts)"]}]}
    assert run(cv=cv) == []


def test_qualified_skills_derived_from_skills_table_items_only():
    rows = ROWS + [
        ["Python (heading cells are not items)", "C"],
        ["Divers", "R (stats), X (y)"],   # 1-char names are ignored
    ]
    assert qualified_skills_from_rows(rows) == {
        "Python": "lecture/adaptation de scripts",
        "Clean Architecture": "usage pragmatique",
    }


def test_any_master_qualified_skill_must_keep_its_qualifier():
    cv = {"skills_sections": [{"heading": "Méthodes", "items": ["clean architecture", "SOLID"]}]}
    errs = [e for e in run(cv=cv) if e.startswith("[calibration]")]
    assert errs == ["[calibration] tailored_cv.json: 'Clean Architecture' must keep the master qualifier '(usage pragmatique)'"]


def test_calibration_rule_is_not_user_specific():
    # Without a qualified-skills map, a bare "Python" is not a calibration error
    # even though the master text happens to contain "python (lecture".
    docs = {"tailored_cv.json": {"skills_sections": [{"heading": "Langages", "items": ["Python"]}]}}
    errs = check(None, MASTER, FAMILIAR, "", docs, {})
    assert not any(e.startswith("[calibration]") for e in errs)


def test_familiar_skill_in_letter_flagged():
    assert any(e.startswith("[familiar]") for e in run(letter={"paragraphs": ["J'utilise Docker."]}))


def test_jd_contagion_flagged():
    errs = run(cv={"summary_paragraphs": ["Expert Azure DevOps"]}, offer={"technologies": ["Azure DevOps", "C#"]})
    assert any(e.startswith("[jd-contagion]") for e in errs)


def test_calibration_bullets_parsed():
    txt = "## Skill calibration (x)\n\nintro\n\n- **WCF** a\n- **WPF** b\n\n## Hidden skills\n- z"
    assert calibration_bullets(txt) == ["**WCF** a", "**WPF** b"]


def test_jd_tech_named_as_an_admitted_gap_is_allowed():
    letter = {"paragraphs": ["Je n'ai pas d'expérience Azure DevOps."]}
    assert run(letter=letter, offer={"technologies": ["Azure DevOps"]}) == []


def _earlier(meta, bullet, extra=None):
    exp = [{"role_line": "Expériences antérieures", "metadata_line": meta, "bullets": [bullet]}]
    if extra:
        exp.insert(0, extra)
    return {"experience": exp}


def test_earlier_line_with_one_clause_per_employer_passes():
    cv = _earlier("Peaktime SAS | JFC Informatique & Média (Asnières) | ROCC Computers Ltd",
                  "Parcours en C et C++. Peaktime : logiciel TV. JFC : radio. ROCC : OS propriétaire.",
                  {"role_line": "X", "metadata_line": "JFC Informatique & Média (Kantar) | Paris | 2002", "bullets": []})
    assert run(cv=cv) == []


def test_earlier_line_missing_clause_flagged():
    cv = _earlier("Peaktime SAS | ROCC Computers Ltd", "Développement C++ et C sous Unix, assembleur.")
    assert any(e.startswith("[earlier-line]") for e in run(cv=cv))


def test_earlier_line_same_name_as_full_entry_must_be_disambiguated():
    cv = _earlier("JFC Informatique & Média", "JFC : radio.",
                  {"role_line": "X", "metadata_line": "JFC Informatique & Média (Kantar) | Paris | 2002", "bullets": []})
    assert any("disambiguator" in e for e in run(cv=cv))


def test_jd_tech_named_as_something_to_learn_is_allowed():
    letter = {"paragraphs": ["La cryptographie reste à apprendre."]}
    assert run(letter=letter, offer={"technologies": ["Cryptographie"]}) == []


def test_jd_tech_in_a_sentence_with_nouveau_is_still_flagged():
    letter = {"paragraphs": ["J'ai déployé Kubernetes sur un nouveau projet."]}
    assert any(e.startswith("[jd-contagion]") for e in run(letter=letter, offer={"technologies": ["Kubernetes"]}))


# ----------------------------------------------------- rule 6, cold flow

PROFILE = {"tech_stack_hints": [
    "Kubernetes, Terraform (CI/CD, en bonus sur l'offre)",
    "C# / .NET — inferred from job listings",
]}


def test_stack_hint_techs_drop_annotations():
    assert stack_hint_techs(PROFILE) == ["Kubernetes", "Terraform", "C#", ".NET"]


def test_stack_contagion_flagged_when_no_offer():
    letter = {"paragraphs": ["J'ai déployé Kubernetes en production."]}
    errs = run(letter=letter, profile=PROFILE)
    assert len(errs) == 1 and errs[0].startswith("[stack-contagion] letter.json: 'Kubernetes'")


def test_stack_tech_that_is_in_master_cv_is_allowed():
    assert run(letter={"paragraphs": ["Dix ans de C#."]}, profile=PROFILE) == []


def test_stack_tech_named_as_a_gap_is_allowed():
    letter = {"paragraphs": ["Je n'ai pas encore utilisé Terraform."]}
    assert run(letter=letter, profile=PROFILE) == []


def test_offer_technologies_take_precedence_over_profile():
    letter = {"paragraphs": ["J'ai déployé Kubernetes en production."]}
    errs = run(letter=letter, offer={"technologies": ["Kubernetes"]}, profile=PROFILE)
    assert len(errs) == 1 and errs[0].startswith("[jd-contagion] letter.json:")


def test_run_check_reads_company_profile_in_cold_prep(tmp_path, monkeypatch):
    """run_check picks up company_profile.json when job_offer_analysis.json is absent."""
    import json

    from docx import Document

    from scripts.claims_guard import run_check

    master = tmp_path / "MASTER_CV.docx"
    doc = Document()
    doc.add_paragraph("Développeur C# senior.")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Langages"
    t.rows[0].cells[1].text = "C#, Python (lecture/adaptation de scripts)"
    doc.save(str(master))
    monkeypatch.setenv("JOB_TAILOR_HOME", str(tmp_path))

    prep = tmp_path / "_prep"
    prep.mkdir()
    (prep / "company_profile.json").write_text(json.dumps(PROFILE), encoding="utf-8")
    (prep / "letter.json").write_text(json.dumps(
        {"paragraphs": ["J'ai déployé Kubernetes en production."]}), encoding="utf-8")
    (prep / "tailored_cv.json").write_text(json.dumps(
        {"skills_sections": [{"heading": "Langages", "items": ["C#", "Python"]}]}), encoding="utf-8")

    errs = run_check(prep, master=master)
    assert any(e.startswith("[stack-contagion] letter.json: 'Kubernetes'") for e in errs)
    assert "[calibration] tailored_cv.json: 'Python' must keep the master qualifier '(lecture/adaptation de scripts)'" in errs


# ------------------------------------------- addendum prohibitions are not a source

def test_addendum_prohibition_does_not_source_the_forbidden_tech():
    cv = {"skills_sections": [{"heading": "Données", "items": ["Entity Framework"]}]}
    errs = run(cv=cv, addendum="Écrire SQL Server, jamais Entity Framework.")
    assert any(e.startswith("[unsourced-skill]") and "Entity Framework" in e for e in errs)


def test_addendum_positive_line_still_sources_a_skill():
    cv = {"skills_sections": [{"heading": "Desktop", "items": ["Delphi"]}]}
    assert run(cv=cv, addendum="- **Delphi — logiciels Telmar Peaktime.**") == []


def test_addendum_mixed_line_keeps_positive_sentence_drops_prohibition():
    addendum = ("- **Delphi — logiciels Telmar.** Ne pas l'étendre à Razor.\n"
                "- Razor, Azure DevOps ne sont pas dans le fact base.")
    cv = {"skills_sections": [{"heading": "Outils", "items": ["Delphi", "Razor", "Azure DevOps"]}]}
    errs = run(cv=cv, addendum=addendum)
    assert sorted(e.split("'")[3] for e in errs) == ["Azure DevOps", "Razor"]
    assert all(e.startswith("[unsourced-skill]") for e in errs)


def test_addendum_prohibition_does_not_exempt_jd_contagion():
    letter = {"paragraphs": ["J'ai utilisé Razor en production."]}
    errs = run(letter=letter, offer={"technologies": ["Razor"]},
               addendum="Razor n'est pas dans le fact base, ne jamais le reprendre.")
    assert any(e.startswith("[jd-contagion] letter.json: 'Razor'") for e in errs)


# ------------------------------------- addendum sections and clause-level negation

from scripts.claims_guard import addendum_source_text, familiar_items  # noqa: E402


def test_no_code_is_not_a_negation_in_the_addendum():
    cv = {"skills_sections": [{"heading": "Low-code", "items": ["Glide"]}]}
    addendum = "## Off-CV facts to remember\n- Comparé trois plateformes no-code, dont Glide."
    assert run(cv=cv, addendum=addendum) == []


def test_hidden_skills_bullets_are_never_filtered():
    bullet = "- Diagnostic d'un portage qui ne pouvait pas être livré sans équivalent VSS fiable."
    cv = {"skills_sections": [{"heading": "Systèmes", "items": ["VSS"]}]}
    assert run(cv=cv, addendum="## Hidden skills\n" + bullet) == []
    errs = run(cv=cv, addendum="## Off-CV facts to remember\n" + bullet)
    assert any(e.startswith("[unsourced-skill]") and "'VSS'" in e for e in errs)


def test_negation_drops_only_its_clause():
    src = addendum_source_text("## Additional experience entries\n### Acme — 2020\n"
                               "- Bases : PostgreSQL ; SQLite en local. Pas de SQL Server ici.")
    assert "postgresql" in src and "sqlite en local." in src
    assert "sql server" not in src


def test_skill_calibration_sentences_with_a_negation_are_dropped_whole():
    addendum = "## Skill calibration\n- **Kafka — évalué, pas retenu.** Go : usage quotidien."
    src = addendum_source_text(addendum)
    assert "kafka" not in src and "go : usage quotidien." in src


def test_jd_tech_in_an_affirmative_clause_of_a_contrast_is_flagged():
    letter = {"paragraphs": ["J'ai déployé Kubernetes, pas Terraform."]}
    errs = run(letter=letter, offer={"technologies": ["Kubernetes", "Terraform"]})
    assert [e.split("'")[1] for e in errs] == ["Kubernetes"]


def test_familiar_items_drop_tool_qualifiers_listed_elsewhere():
    rows = [["Outils", "GitHub, GitLab, Jira"],
            ["Familiar / to deepen", "Docker, GitHub / GitLab CI/CD, RabbitMQ / Kafka"]]
    assert familiar_items(rows) == ["Docker", "CI/CD", "RabbitMQ", "Kafka"]


# --------------------------------------------------- earlier-experience line

def test_earlier_line_accepts_english_colon_and_legal_suffixes():
    cv = _earlier("Globex SA | Initech GmbH", "Globex SA: built tools. Initech: payroll in C.")
    assert run(cv=cv) == []
    cv = _earlier("Globex SA | Initech GmbH", "Initech: payroll. Globex: built tools.")
    assert any("out of order" in e for e in run(cv=cv))


def test_earlier_line_messages_name_the_file_and_are_generic():
    cv = _earlier("Globex SA", "Built tools.")
    errs = run(cv=cv)
    assert errs == ["[earlier-line] tailored_cv.json: no 'Employer: ...' clause for 'Globex SA' "
                    "in the earlier-experience bullet"]
    cv = _earlier("Globex SA", "Globex : x.",
                  {"role_line": "X", "metadata_line": "Globex SA | Lyon | 2010", "bullets": []})
    assert any("'(<city>)'" in e for e in run(cv=cv))


# ------------------------------------------------------------- run_check(only=)

def _prep_with_cv_and_letter_problems(tmp_path, monkeypatch):
    import json

    from docx import Document

    master = tmp_path / "MASTER_CV.docx"
    doc = Document()
    doc.add_paragraph("Développeur C# senior.")
    doc.save(str(master))
    monkeypatch.setenv("JOB_TAILOR_HOME", str(tmp_path))
    prep = tmp_path / "_prep"
    prep.mkdir()
    (prep / "job_offer_analysis.json").write_text(json.dumps({"technologies": ["Kubernetes"]}), encoding="utf-8")
    (prep / "tailored_cv.json").write_text(json.dumps(
        {"skills_sections": [{"heading": "Langages", "items": ["C#", "Oracle"]}]}), encoding="utf-8")
    (prep / "letter.json").write_text(json.dumps({"paragraphs": ["J'ai déployé Kubernetes."]}), encoding="utf-8")
    (prep / "linkedin.json").write_text(json.dumps({"connection_request": "Bonjour"}), encoding="utf-8")
    return prep, master


def test_run_check_only_scopes_to_the_named_files(tmp_path, monkeypatch):
    from scripts.claims_guard import run_check

    prep, master = _prep_with_cv_and_letter_problems(tmp_path, monkeypatch)
    every = run_check(prep, master=master)
    assert any("tailored_cv.json:" in e for e in every) and any("letter.json:" in e for e in every)
    cv_only = run_check(prep, master=master, only=("tailored_cv.json",))
    assert cv_only and all("tailored_cv.json:" in e for e in cv_only)
    letters = run_check(prep, master=master, only=("letter.json", "short_letter.json"))
    assert letters == [e for e in every if " letter.json:" in e]
    assert run_check(prep, master=master, only=("linkedin.json",)) == []
