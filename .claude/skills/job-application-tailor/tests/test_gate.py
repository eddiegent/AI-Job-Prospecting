"""Tests for `cli.py gate <stage> <folder>` and `cli.py read-cv`
(scripts/commands/gate.py, scripts/commands/pipeline.py)."""
from __future__ import annotations

import copy
import json
from argparse import Namespace
from pathlib import Path

import pytest

from scripts.commands.gate import cmd_gate
from scripts.commands.pipeline import cmd_read_cv
from tests.conftest import load_fixture


@pytest.fixture
def home(tmp_path, monkeypatch) -> Path:
    """User data dir with a tiny master CV and no prefs/addendum."""
    from docx import Document

    data = tmp_path / "userdata"
    data.mkdir()
    doc = Document()
    doc.add_paragraph("Alex Dupont")
    doc.add_paragraph("Senior Engineer, C++ and API platforms.", style="Heading 1")
    doc.add_paragraph("   ")  # blank paragraphs are skipped by read-cv
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text = "Langages"
    t.rows[0].cells[1].text = "C++, Python (scripts)"
    t.rows[1].cells[0].text = "Familier / a approfondir"
    t.rows[1].cells[1].text = "Docker"
    doc.save(str(data / "MASTER_CV.docx"))
    monkeypatch.setenv("JOB_TAILOR_HOME", str(data))
    return data


def _run(tmp_path: Path, files: dict[str, dict]) -> Path:
    run = tmp_path / "good-14072026-role"
    prep = run / "_prep"
    prep.mkdir(parents=True)
    for name, payload in files.items():
        (prep / name).write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return run


def _gate(stage: str, folder: Path, capsys) -> tuple[int, str]:
    with pytest.raises(SystemExit) as e:
        cmd_gate(None, Namespace(stage=stage, folder=str(folder), cv=None))
    return e.value.code, capsys.readouterr().out


def _offer() -> dict:
    return {"job_title": "Dev", "company_name": "Acme", "required_skills": ["C++"],
            "preferred_skills": [], "responsibilities": [], "technologies": ["Kubernetes", "C++"],
            "ats_keywords": []}


def test_gate_unknown_folder_exits_2(tmp_path, home):
    with pytest.raises(SystemExit) as e:
        cmd_gate(None, Namespace(stage="cv", folder=str(tmp_path / "nope"), cv=None))
    assert e.value.code == 2


def test_gate_missing_input_fails_and_stops(tmp_path, home, capsys):
    code, out = _gate("roles", _run(tmp_path, {}), capsys)
    assert code == 1
    assert "[FAIL] schema role_candidates.json" in out
    assert "role_candidates.json not found" in out
    assert "role grounding" not in out  # later checks skipped
    assert "skipped" in out


def test_gate_match_recounts_in_place_and_passes(tmp_path, home, capsys):
    ma = load_fixture("sample_match_analysis.json")
    ma["match_summary"]["overall_fit_pct"] = 12  # LLM drift
    run = _run(tmp_path, {"match_analysis.json": ma, "job_offer_analysis.json": _offer(),
                          "cv_fact_base_merged.json": load_fixture("sample_fact_base.json")})
    code, out = _gate("match", run, capsys)
    assert code == 0, out
    assert "[PASS] recount match_summary: updated" in out
    assert "[PASS] schema match_analysis.json" in out
    assert "[PASS] match grounding (vs cv_fact_base_merged.json)" in out
    saved = json.loads((run / "_prep" / "match_analysis.json").read_text(encoding="utf-8"))
    assert saved["match_summary"]["overall_fit_pct"] != 12
    # second run: already in sync
    code, out = _gate("match", run, capsys)
    assert "[PASS] recount match_summary: unchanged" in out


def _direct(req: str) -> dict:
    return {"match_summary": {}, "matches": [
        {"requirement": req, "category": "technology", "match_type": "direct", "evidence": "x"}]}


def test_gate_match_fails_without_merged_fact_base(tmp_path, home, capsys):
    """No silent fallback to cv_fact_base.json: the addendum would be ignored."""
    run = _run(tmp_path, {"match_analysis.json": _direct("C++"), "job_offer_analysis.json": _offer(),
                          "cv_fact_base.json": load_fixture("sample_fact_base.json")})
    code, out = _gate("match", run, capsys)
    assert code == 1
    assert "[FAIL] match grounding (vs cv_fact_base_merged.json)" in out
    assert "cv_fact_base_merged.json not found" in out
    assert "python scripts/claims_guard.py merge" in out and "Step 3.9" in out


def test_gate_match_grounds_against_merged_fact_base(tmp_path, home, capsys):
    fb = load_fixture("sample_fact_base.json")
    run = _run(tmp_path, {"match_analysis.json": _direct("Kubernetes"), "job_offer_analysis.json": _offer(),
                          "cv_fact_base_merged.json": fb})
    code, out = _gate("match", run, capsys)
    assert code == 1 and "[FAIL] match grounding (vs cv_fact_base_merged.json)" in out

    merged = copy.deepcopy(fb)
    merged["addendum_hidden_skills"] = ["Kubernetes"]
    (run / "_prep" / "cv_fact_base_merged.json").write_text(json.dumps(merged), encoding="utf-8")
    code, out = _gate("match", run, capsys)
    assert code == 0, out
    assert "[PASS] match grounding (vs cv_fact_base_merged.json)" in out


@pytest.mark.parametrize("detail, passes", [
    (None, False),                                         # no evidence at all
    ("Pas de Kubernetes à ce poste.", False),              # negated addendum line is not evidence
    ("No Kubernetes in this role, only Docker.", False),
    ("Déploiement des services sur Kubernetes.", True),    # positive line grounds it
    ("Pas de Docker ; déploiement sur Kubernetes.", True),  # only the negated clause is dropped
])
def test_gate_match_negated_addendum_detail_does_not_ground(tmp_path, home, capsys, detail, passes):
    merged = load_fixture("sample_fact_base.json")
    if detail:
        merged["experience"][0]["details"].append(detail)  # where merge_addendum_into_fact_base puts it
    run = _run(tmp_path, {"match_analysis.json": _direct("Kubernetes"), "job_offer_analysis.json": _offer(),
                          "cv_fact_base_merged.json": merged})
    code, out = _gate("match", run, capsys)
    assert (code == 0) is passes, out


@pytest.mark.parametrize("stage, target", [("roles", "role_candidates.json"),
                                           ("selected-role", "selected_role.json")])
def test_gate_role_grounding_fails_without_merged_fact_base(tmp_path, home, capsys, stage, target):
    role = {"title": "Senior Engineer", "source": "generalist", "seniority_band": "senior",
            "generated_at": "2026-10-20T10:00:00Z", "company_name": "Acme", "language": "en",
            "emphasis_areas": ["C++"], "rationale": "x"}
    payload = ({"candidates": [{k: role[k] for k in ("title", "rationale", "seniority_band", "emphasis_areas")}],
                "generated_at": role["generated_at"], "company_name": "Acme"}
               if stage == "roles" else role)
    run = _run(tmp_path, {target: payload, "company_profile.json": {"tech_stack_hints": ["Kubernetes"]},
                          "cv_fact_base.json": load_fixture("sample_fact_base.json")})
    code, out = _gate(stage, run, capsys)
    assert code == 1, out
    assert "cv_fact_base_merged.json not found" in out and "claims_guard.py merge" in out


def test_gate_cv_reports_every_failure_together(tmp_path, home, capsys):
    (home / "user_prefs.yaml").write_text('forbidden_title_labels:\n  - "Backend"\n', encoding="utf-8")
    cv = load_fixture("tailored_cv_good.json")
    cv["title"] = "Backend Engineer"
    cv["skills_sections"] = [{"heading": "Langages", "items": ["C++", "Python"]}]
    code, out = _gate("cv", _run(tmp_path, {"tailored_cv.json": cv}), capsys)
    assert code == 1
    assert "[PASS] schema tailored_cv.json" in out
    assert "[FAIL] forbidden title labels in tailored_cv.json" in out
    assert "[FAIL] claims check (tailored_cv.json)" in out
    assert "tailored_cv.json: 'Python' must keep the master qualifier '(scripts)'" in out
    assert "gate cv: FAIL (2 check(s) failed)" in out


def test_gate_cv_schema_failure_stops_stage(tmp_path, home, capsys):
    code, out = _gate("cv", _run(tmp_path, {"tailored_cv.json": {"title": "x"}}), capsys)
    assert code == 1
    assert "[FAIL] schema tailored_cv.json" in out
    assert "claims check" not in out


def _clean_cv() -> dict:
    cv = load_fixture("tailored_cv_good.json")
    cv["experience"][-1]["bullets"] = ["TinyShop : premiers postes en développement logiciel (C++)."]
    return cv


def _bad_cv() -> dict:
    cv = _clean_cv()
    cv["skills_sections"] = [{"heading": "Langages", "items": ["C++", "Python"]}]  # loses its qualifier
    return cv


def _linkedin(msg: str = "Bonjour, échange possible ?") -> dict:
    return {"variants": [{"target": "recruiter", "message": msg}]}


def _letter(par: str) -> dict:
    return {"greeting": "Bonjour,", "paragraphs": [par], "signoff": "Cordialement,", "name": "Alex Dupont"}


def test_gate_linkedin_ignores_problems_in_other_stages_files(tmp_path, home, capsys):
    run = _run(tmp_path, {"tailored_cv.json": _bad_cv(), "linkedin.json": _linkedin(),
                          "letter.json": _letter("J'utilise Docker au travail.")})
    code, out = _gate("linkedin", run, capsys)
    assert code == 0, out
    assert "[PASS] claims check (linkedin.json)" in out


def test_gate_cv_ignores_letter_problems_and_letters_ignore_cv(tmp_path, home, capsys):
    files = {"tailored_cv.json": _clean_cv(),
             "letter.json": _letter("J'utilise Docker au travail."), "short_letter.json": _letter("Bonjour.")}
    run = _run(tmp_path, files)
    code, out = _gate("cv", run, capsys)
    assert code == 0, out
    assert "[PASS] claims check (tailored_cv.json)" in out

    files["tailored_cv.json"] = _bad_cv()
    files["letter.json"] = _letter("Bonjour.")
    run2 = _run(tmp_path / "b", files)
    code, out = _gate("letters", run2, capsys)
    assert code == 0, out
    assert "[PASS] claims check (letter.json, short_letter.json)" in out


def test_gate_letters_messages_name_the_file(tmp_path, home, capsys):
    offer = _offer()
    run = _run(tmp_path, {"job_offer_analysis.json": offer,
                          "letter.json": _letter("J'ai déployé Kubernetes en production."),
                          "short_letter.json": _letter("J'utilise Docker au travail.")})
    code, out = _gate("letters", run, capsys)
    assert code == 1
    assert "[jd-contagion] letter.json: 'Kubernetes'" in out
    assert "[familiar] short_letter.json: mentions familiar-only skill 'Docker'" in out


def test_read_cv_prints_paragraphs_then_tables(home, capsys):
    cmd_read_cv(None, Namespace(cv=None))
    assert capsys.readouterr().out.splitlines() == [
        "[Normal] Alex Dupont",
        "[Heading 1] Senior Engineer, C++ and API platforms.",
        "--- TABLE 0 ---",
        "Langages | C++, Python (scripts)",
        "Familier / a approfondir | Docker",
    ]


def test_read_cv_missing_cv_exits_2(tmp_path):
    with pytest.raises(SystemExit) as e:
        cmd_read_cv(None, Namespace(cv=str(tmp_path / "missing.docx")))
    assert e.value.code == 2
